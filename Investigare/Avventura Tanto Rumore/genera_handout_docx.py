"""Raccoglie i 13 handout in un DOCX di 13 pagine pronto per la stampa.

Le pagine sono incorporate a 300 dpi per conservare grafica e caratteri giapponesi.
I contenuti modificabili restano nel wizard e negli HTML esportati.
"""
from pathlib import Path
import concurrent.futures
import io
import json
import subprocess
import tempfile

import fitz
from bs4 import BeautifulSoup
from docx import Document
from docx.shared import Mm, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
OUT = ROOT / 'storico/handout_docx_2026-09-19'
CHROME = r'C:\Program Files\Google\Chrome\Application\chrome.exe'
CSS = '''
@page { size: A4; margin: 0; }
@media print {
  html, body { margin: 0; padding: 0; background: white; }
  .ho-foglio { width: 210mm; min-height: 0; padding: 12mm; margin: 0; box-shadow: none; }
  .ho-tabulato { line-height: 1.18; }
  .ho-tabulato .ho-intestazione { margin-bottom: 7px; padding-bottom: 5px; }
  .ho-tabulato .ho-meta { margin-bottom: 7px; }
  .ho-tabulato td, .ho-tabulato th { padding: 2px 4px; line-height: 1.15; }
  .ho-doc img { display: inline-block; vertical-align: top; max-width: 48%; max-height: 135mm; width: auto; height: auto; margin: 4mm 1% 0 0; }
}
'''

def render(source):
    original = source.read_text(encoding='utf-8')
    assert 'ho-avviso-gm' not in original
    soup = BeautifulSoup(original, 'html.parser')
    base = soup.new_tag('base', href=source.parent.as_uri() + '/')
    soup.head.insert(0, base)
    style = soup.new_tag('style')
    soup.head.append(style)
    html_path = OUT / source.name
    pdf_path = OUT / (source.stem + '.pdf')
    # Fit complete handouts; never crop, remove text, or split one over two pages.
    for scale in (1, .96, .92, .88, .84):
        style.string = CSS + f'\n@media print {{ .ho-foglio {{ zoom: {scale}; width: {210 / scale}mm; }} }}'
        html_path.write_text(str(soup), encoding='utf-8')
        with tempfile.TemporaryDirectory(prefix='bakuon-docx-') as profile:
            assert Path(profile).resolve().is_relative_to(Path(tempfile.gettempdir()).resolve())
            result = subprocess.run([CHROME, '--headless', '--disable-gpu', '--no-first-run',
                '--no-default-browser-check', '--disable-background-networking', '--disable-extensions',
                '--no-pdf-header-footer', '--user-data-dir=' + profile,
                '--print-to-pdf=' + str(pdf_path), html_path.as_uri()],
                capture_output=True, timeout=45, creationflags=0x08000000)
            assert result.returncode == 0 and pdf_path.exists(), source.name
        with fitz.open(pdf_path) as pdf:
            if len(pdf) != 1:
                continue
            page = pdf[0]
            assert 'AVVISO PER IL MASTER' not in page.get_text()
            for block in page.get_text('dict')['blocks']:
                for line in block.get('lines', []):
                    for span in line['spans']:
                        x0, y0, x1, y1 = span['bbox']
                        assert 0 <= x0 < x1 <= page.rect.width and 0 <= y0 < y1 <= page.rect.height, source.name
            png = page.get_pixmap(dpi=300).tobytes('png')
            page.get_pixmap(dpi=100).save(OUT / (source.stem + '_anteprima.png'))
            return source.name, png, scale
    raise RuntimeError('Non entra su una pagina: ' + source.name)

def main():
    OUT.mkdir(exist_ok=True)
    sources = sorted((ROOT / 'handout').glob('*.html'))
    assert len(sources) == 13
    with concurrent.futures.ThreadPoolExecutor(max_workers=3) as pool:
        pages = list(pool.map(render, sources))
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Mm(5)
    sec.header_distance = sec.footer_distance = Mm(0)
    normal = doc.styles['Normal']
    normal.font.size = Pt(1)
    normal.paragraph_format.space_before = normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1
    for i, (name, png, scale) in enumerate(pages):
        p = doc.add_paragraph()
        p.paragraph_format.page_break_before = i > 0
        p.paragraph_format.widow_control = False
        picture = p.add_run().add_picture(io.BytesIO(png), width=Mm(200))
        picture._inline.docPr.set('descr', name)
        print(f'{i + 1:02d}: {name} - una pagina, scala {scale}', flush=True)
    settings = doc.settings.element
    no_compress = OxmlElement('w:doNotAutoCompressPictures')
    settings.append(no_compress)
    doc.core_properties.title = 'Bakuon - Handout da stampare'
    doc.core_properties.subject = '13 handout, uno per pagina, senza note GM'
    doc.core_properties.author = 'Maru'
    target = ROOT / 'HANDOUT_STAMPA.docx'
    doc.save(target)
    (OUT / 'manifest.json').write_text(json.dumps([{'file': n, 'scala': s} for n, _, s in pages], indent=2), encoding='utf-8')
    print(target)

if __name__ == '__main__':
    main()
